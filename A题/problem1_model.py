"""问题1：无限长圆柱药材的径向湿热传导模型。

本文件只建立连续模型及其径向有限体积离散形式，不进行数值求解、绘图或结果导出。
"""
from pathlib import Path
import numpy as np
from openpyxl import load_workbook

RADIUS = 0.02
RHO, CP, K = 820.0, 2600.0, 0.36
H_T, H_M = 25.0, 8.0e-7
T0, C0 = 28.0, 2.55


def moisture_diffusivity(C):
    """D(C)，单位 m^2/s。"""
    return 7.0e-9 * np.exp(-0.89 / np.maximum(C, 1.0e-12))


def load_boundary_data(path=None):
    """读取附件1中的时间、烘房温度和烘房水分浓度。"""
    path = path or Path(__file__).resolve().parent / "附件" / "附件1.xlsx"
    wb = load_workbook(path, data_only=True, read_only=True)
    rows = np.asarray(list(wb.active.iter_rows(min_row=2, values_only=True)), dtype=float)
    wb.close()
    if rows.ndim != 2 or rows.shape[1] < 3:
        raise ValueError("附件1.xlsx应包含时间、温度和水分浓度三列")
    return rows[:, 0], rows[:, 1], rows[:, 2]


def radial_grid(n=200):
    """返回单元中心、单元面和圆柱径向体积因子。"""
    dr = RADIUS / n
    faces = np.arange(n + 1) * dr
    centers = (faces[:-1] + faces[1:]) / 2
    volumes = 0.5 * (faces[1:] ** 2 - faces[:-1] ** 2)
    return centers, faces, volumes


def semi_discrete_model(t, temperature, concentration, grid, boundary):
    """返回有限体积离散后的温度场和水分场变化率。"""
    _, faces, volumes = grid
    time_air, temp_air, moist_air = boundary
    n = len(temperature)
    dr = faces[1] - faces[0]
    t_air = np.interp(t, time_air, temp_air)
    c_air = np.interp(t, time_air, moist_air)
    heat_flux = np.zeros(n + 1)
    moisture_flux = np.zeros(n + 1)
    heat_flux[1:n] = -K * np.diff(temperature) / dr
    d_face = moisture_diffusivity(0.5 * (concentration[:-1] + concentration[1:]))
    moisture_flux[1:n] = -d_face * np.diff(concentration) / dr
    heat_flux[n] = H_T * (temperature[-1] - t_air)
    moisture_flux[n] = H_M * (concentration[-1] - c_air)
    heat_div = (faces[1:] * heat_flux[1:] - faces[:-1] * heat_flux[:-1]) / volumes
    moisture_div = (faces[1:] * moisture_flux[1:] - faces[:-1] * moisture_flux[:-1]) / volumes
    return -heat_div / (RHO * CP), -moisture_div


def initial_state(n=200):
    """返回均匀初始温度场和水分浓度场。"""
    return np.full(n, T0), np.full(n, C0)
