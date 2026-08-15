from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "\u6298\u7ebf\u56fe"
OUT.mkdir(exist_ok=True)
plt.rcParams.update({
    "font.sans-serif": ["Microsoft YaHei", "SimHei", "DejaVu Sans"],
    "axes.unicode_minus": False,
    "font.size": 11,
    "pdf.fonttype": 42,
})


def read(number, columns):
    path = ROOT / "\u9644\u4ef6" / f"\u9644\u4ef6{number}.xlsx"
    workbook = load_workbook(path, data_only=True, read_only=True)
    rows = list(workbook.active.iter_rows(min_row=2, values_only=True))
    workbook.close()
    data = np.array(rows, dtype=float)
    assert data.shape[1] == columns and np.isfinite(data).all(), path
    assert np.all(np.diff(data[:, 0]) > 0), "Time must increase"
    print(f"Attachment {number}: {len(data)} points, {data[0, 0]:g}-{data[-1, 0]:g} s")
    return data


def draw(ax, x, y, title, ylabel, color):
    ax.plot(x, y, color=color, linewidth=1.7)
    ax.set(title=title, xlabel="\u65f6\u95f4 / s", ylabel=ylabel)
    ax.grid(True, color="#dddddd", linewidth=0.6, alpha=0.8)
    ax.set_axisbelow(True)
    ax.spines[["top", "right"]].set_visible(False)
    ax.set_xlim(x[0], x[-1])
    ax.ticklabel_format(axis="x", style="plain", useOffset=False)


def save(fig, name):
    for extension in ("png", "pdf"):
        fig.savefig(OUT / f"{name}.{extension}", dpi=300, bbox_inches="tight")
    plt.close(fig)


def main():
    first = read(1, 3)
    second = read(2, 2)
    fig, axes = plt.subplots(2, 1, figsize=(10, 7.5), layout="constrained")
    draw(axes[0], first[:, 0], first[:, 1],
         "\u9644\u4ef6\u4e00\uff1a\u70d8\u623f\u6e29\u5ea6\u968f\u65f6\u95f4\u7684\u53d8\u5316",
         "\u6e29\u5ea6 / \u00b0C", "#c34343")
    draw(axes[1], first[:, 0], first[:, 2],
         "\u9644\u4ef6\u4e00\uff1a\u70d8\u623f\u6c34\u5206\u6d53\u5ea6\u968f\u65f6\u95f4\u7684\u53d8\u5316",
         "\u6c34\u5206\u6d53\u5ea6 / (kg/kg)", "#2477a8")
    save(fig, "\u9644\u4ef61_\u6e29\u5ea6\u4e0e\u6c34\u5206\u6d53\u5ea6\u6298\u7ebf\u56fe")
    fig, ax = plt.subplots(figsize=(10, 4.8), layout="constrained")
    draw(ax, second[:, 0], second[:, 1],
         "\u9644\u4ef6\u4e8c\uff1a\u836f\u6750\u534a\u5f84\u968f\u65f6\u95f4\u7684\u53d8\u5316",
         "\u534a\u5f84 / cm", "#25836b")
    save(fig, "\u9644\u4ef62_\u534a\u5f84\u6298\u7ebf\u56fe")


if __name__ == "__main__":
    main()
