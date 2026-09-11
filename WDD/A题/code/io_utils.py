# -*- coding: utf-8 -*-
"""数据读写：附件、边界驱动、结果 xlsx（照附件 3 模板格式）。"""
import numpy as np
import openpyxl
from openpyxl import Workbook

ATT1 = '附件/附件1.xlsx'
ATT2 = '附件/附件2.xlsx'


def load_attachment1(path=ATT1):
    """返回 (t[s], T[°C], C[kg/kg])，241 个点，步长 60 s。"""
    ws = openpyxl.load_workbook(path).active
    rows = list(ws.iter_rows(values_only=True))[1:]
    t = np.array([r[0] for r in rows], float)
    T = np.array([r[1] for r in rows], float)
    C = np.array([r[2] for r in rows], float)
    return t, T, C


def load_attachment2(path=ATT2):
    """返回 (t[s], R[cm])，145 个点，步长 1800 s。"""
    ws = openpyxl.load_workbook(path).active
    rows = list(ws.iter_rows(values_only=True))[1:]
    t = np.array([r[0] for r in rows], float)
    R = np.array([r[1] for r in rows], float)
    return t, R


def make_boundary(t, T, C):
    """边界驱动函数：数据内线性插值，超出后平台常值延拓（HC-21）。"""
    def bc(tt):
        if tt <= t[-1]:
            return float(np.interp(tt, t, T)), float(np.interp(tt, t, C))
        return float(T[-1]), float(C[-1])
    return bc


def make_radius(t, R_cm):
    """R(t) 与 Rdot(t)：PCHIP 保形单调插值（无过冲），返回 SI 单位（m, m/s）。"""
    from scipy.interpolate import PchipInterpolator
    p = PchipInterpolator(t, R_cm / 100.0)
    pd = p.derivative()
    def R_func(tt):
        return float(p(min(tt, t[-1])))
    def Rdot_func(tt):
        return float(pd(min(tt, t[-1])))
    return R_func, Rdot_func


def _write_sheet(ws, title, t_arr, r_arr_cm, data, surface_col=False):
    """按模板写表：A1 '时间\\到药材中心的距离'，第 1 行距离(cm)，A 列时间(s)。"""
    ws.title = title
    header = ['时间\\到药材中心的距离'] + [round(float(r), 4) for r in r_arr_cm]
    if surface_col:
        header[-1] = '药材表面'
    ws.append(header)
    for i, tt in enumerate(t_arr):
        row = [round(float(tt), 4)]
        for v in data[i]:
            row.append(None if (v is None or (isinstance(v, float) and np.isnan(v)))
                       else round(float(v), 4))
        ws.append(row)


def save_result(path, sheets):
    """sheets: [(sheet_name, t_arr, r_arr_cm, data, surface_col), ...]

    Windows 下目标文件被 Excel 占用时 PermissionError 高发：
    重试若干次并给出明确提示；openpyxl 清理临时文件失败的 atexit 报错同理。
    """
    import time
    wb = Workbook(write_only=True)
    for name, t_arr, r_arr_cm, data, surface_col in sheets:
        ws = wb.create_sheet()
        _write_sheet(ws, name, t_arr, r_arr_cm, data, surface_col)
    for attempt in range(5):
        try:
            wb.save(path)
            print('已写出', path)
            return
        except PermissionError:
            if attempt == 4:
                raise SystemExit(
                    f'\n[保存失败] {path} 被其他程序占用（通常是 Excel 打开着该文件）。\n'
                    f'请关闭 Excel 中的此文件后重新运行脚本。')
            print(f'  {path} 被占用，2 s 后重试（{attempt + 1}/5）...')
            time.sleep(2)
