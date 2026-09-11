# -*- coding: utf-8 -*-
"""result*.xlsx 交付写盘（附件 3 模板口径）。

模板结构：A1 为 '时间\\到药材中心的距离'，第 1 行为距离（cm），A 列为时间（s）。
result1/2 双工作表（温度、水分浓度），步长 1 s；result3/4 单表，步长 60 s。
数值按 R8 取 round(v,4) 真实写盘（不是显示层四舍五入），同时设 0.0000 数字格式。
"""
from __future__ import annotations

import openpyxl
from openpyxl.cell import WriteOnlyCell

import params as P

CORNER_LABEL = "时间\\到药材中心的距离"
NUMBER_FORMAT = "0.0000"
TIME_FORMAT = "0"
SURFACE_LABEL = P.Q4_TABLE_LAST_COL


def _round4(value):
    return None if value is None else round(float(value), P.DECIMALS)


def _sheet_rows(ws, times_s, radii_labels, field_rows):
    header = [WriteOnlyCell(ws, value=CORNER_LABEL)]
    for label in radii_labels:
        cell = WriteOnlyCell(ws, value=label)
        if isinstance(label, (int, float)):
            cell.number_format = TIME_FORMAT if float(label).is_integer() else "0.0"
        header.append(cell)
    ws.append(header)
    for t, row_vals in zip(times_s, field_rows):
        cells = [WriteOnlyCell(ws, value=int(round(t)))]
        cells[0].number_format = TIME_FORMAT
        for value in row_vals:
            cell = WriteOnlyCell(ws, value=_round4(value))
            cell.number_format = NUMBER_FORMAT
            cells.append(cell)
        ws.append(cells)


def write_two_sheet_result(path, times_s, radii_cm, temperature_rows, moisture_rows,
                           sheet_names=None):
    """result1/result2：温度与水分浓度两个工作表，同一时间/距离网格。"""
    sheet_names = list(P.Q1_SHEETS) if sheet_names is None else list(sheet_names)
    wb = openpyxl.Workbook(write_only=True)
    for name, rows in zip(sheet_names, (temperature_rows, moisture_rows)):
        ws = wb.create_sheet(title=name)
        _sheet_rows(ws, times_s, radii_cm, rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_single_sheet_result(path, times_s, radii_labels, moisture_rows,
                             sheet_name="Sheet1"):
    """result3/result4：单工作表保存水分浓度；域外单元格写空（None）。"""
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet(title=sheet_name)
    _sheet_rows(ws, times_s, radii_labels, moisture_rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def verify_written(path, expect_sheets, expect_dt_s, expect_n_cols, sample_limit=200):
    """回读校验：工作表名、时间步长、列数、四位小数（B-08/B-28/B-29）。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if list(wb.sheetnames) != list(expect_sheets):
        raise AssertionError(f"{path.name} 工作表 {wb.sheetnames} != {expect_sheets}")
    report = {"path": str(path), "sheets": list(wb.sheetnames)}
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = next(rows)
        if header[0] != CORNER_LABEL:
            raise AssertionError(f"{path.name}/{ws.title} A1 != 模板角标")
        if len(header) != expect_n_cols + 1:
            raise AssertionError(f"{path.name}/{ws.title} 列数 {len(header)} != {expect_n_cols + 1}")
        times, checked, bad = [], 0, []
        for row in rows:
            times.append(row[0])
            for value in row[1:]:
                if value is None or checked >= sample_limit:
                    continue
                checked += 1
                if round(float(value), P.DECIMALS) != float(value):
                    bad.append(value)
                if len(f"{float(value):.4f}".split(".")[1]) != P.DECIMALS:
                    bad.append(value)
        if bad:
            raise AssertionError(f"{path.name}/{ws.title} 非四位小数单元格样例 {bad[:5]}")
        if len(times) < 2:
            raise AssertionError(f"{path.name}/{ws.title} 数据行不足")
        step = times[1] - times[0]
        if abs(step - expect_dt_s) > 1e-9:
            raise AssertionError(f"{path.name}/{ws.title} 时间步长 {step} != {expect_dt_s}")
        report[ws.title] = {"n_rows": len(times), "t_first": times[0], "t_last": times[-1],
                            "dt_s": step, "n_cols": len(header) - 1, "sampled_cells": checked}
    wb.close()
    return report
