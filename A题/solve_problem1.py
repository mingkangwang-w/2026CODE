"""用附件1的拟合边界函数计算问题1的表1、表2。

药材视为无限长圆柱，只考虑径向传热和传质。附件1的离散数据用三次样条
拟合为连续的烘房温度 T_inf(t) 和水分浓度 C_inf(t)，作为表面边界条件。
"""
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from scipy.integrate import solve_ivp
from scipy.interpolate import CubicSpline

from problem1_model import (
    C0, CP, H_M, H_T, K, RADIUS, RHO, T0, moisture_diffusivity,
)

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "附件" / "附件1.xlsx"
FULL_RESULT = ROOT / "result1_model.xlsx"
TABLE_RESULT = ROOT / "表1_表2.xlsx"
N = 200


def read_fitted_boundary():
    wb = load_workbook(DATA, data_only=True, read_only=True)
    data = np.asarray(list(wb.active.iter_rows(min_row=2, values_only=True)), dtype=float)
    wb.close()
    t, temperature, concentration = data[:, 0], data[:, 1], data[:, 2]
    # CubicSpline is a piecewise cubic fitting function passing through all measurements.
    return t, CubicSpline(t, temperature), CubicSpline(t, concentration)


def solve():
    t_data, t_inf, c_inf = read_fitted_boundary()
    dr = RADIUS / N
    faces = np.arange(N + 1) * dr
    centers = (faces[:-1] + faces[1:]) / 2
    volumes = 0.5 * (faces[1:] ** 2 - faces[:-1] ** 2)

    def rhs(t, state):
        temperature, concentration = state[:N], state[N:]
        heat_flux = np.zeros(N + 1)
        moisture_flux = np.zeros(N + 1)
        heat_flux[1:N] = -K * np.diff(temperature) / dr
        d_face = moisture_diffusivity(0.5 * (concentration[:-1] + concentration[1:]))
        moisture_flux[1:N] = -d_face * np.diff(concentration) / dr
        heat_flux[N] = H_T * (temperature[-1] - float(t_inf(t)))
        moisture_flux[N] = H_M * (concentration[-1] - float(c_inf(t)))
        heat_div = (faces[1:] * heat_flux[1:] - faces[:-1] * heat_flux[:-1]) / volumes
        moisture_div = (faces[1:] * moisture_flux[1:] - faces[:-1] * moisture_flux[:-1]) / volumes
        return np.r_[-heat_div / (RHO * CP), -moisture_div]

    times = np.arange(0.0, 1800.0 + 1.0, 1.0)
    initial = np.r_[np.full(N, T0), np.full(N, C0)]
    solution = solve_ivp(rhs, (times[0], times[-1]), initial, method="BDF", t_eval=times,
                         rtol=2e-5, atol=1e-8, max_step=10.0)
    if not solution.success:
        raise RuntimeError(solution.message)
    # Add the axis and surface values so requested r=0 and r=R are represented.
    grid_r = np.r_[0.0, centers, RADIUS]
    temperature = np.column_stack((solution.y[0], solution.y[:N].T, solution.y[N - 1]))
    concentration = np.column_stack((solution.y[N], solution.y[N:].T, solution.y[-1]))
    return times, grid_r, temperature, concentration, t_data, t_inf, c_inf


def save_results(times, radii, temperature, concentration):
    def write_sheet(ws, values):
        ws.cell(1, 1, "时间/s")
        for j, radius in enumerate(np.arange(0, RADIUS + 1e-12, 0.001), 2):
            ws.cell(1, j, round(radius * 100, 1))
        for i, time in enumerate(times, 2):
            ws.cell(i, 1, int(time))
            values_at_r = np.array([np.interp(radius, radii, values[i - 2])
                                    for radius in np.arange(0, RADIUS + 1e-12, 0.001)])
            for j, value in enumerate(values_at_r, 2):
                ws.cell(i, j, round(float(value), 4))

    wb = Workbook()
    ws_t = wb.active
    ws_t.title = "温度"
    ws_c = wb.create_sheet("水分浓度")
    write_sheet(ws_t, temperature)
    write_sheet(ws_c, concentration)
    wb.save(FULL_RESULT)

    selected_times = [100, 300, 600, 900, 1200, 1500, 1800]
    selected_radii = np.arange(0, RADIUS + 1e-12, 0.005)
    table_wb = Workbook()
    table_t = table_wb.active
    table_t.title = "表1 温度"
    table_c = table_wb.create_sheet("表2 水分浓度")
    for ws, values in ((table_t, temperature), (table_c, concentration)):
        ws.append(["时间/s"] + [f"{r * 100:g}" for r in selected_radii])
        for time in selected_times:
            i = int(time)
            row = [time] + [round(float(np.interp(r, radii, values[i])), 4) for r in selected_radii]
            ws.append(row)
    table_wb.save(TABLE_RESULT)


def main():
    result = solve()
    save_results(*result[:4])
    _, _, _, _, t_data, t_inf, c_inf = result
    print(f"温度拟合函数：三次样条，数据范围 {t_data[0]:g}-{t_data[-1]:g} s")
    print(f"边界拟合值 t=1800 s：T_inf={float(t_inf(1800)):.4f} °C, C_inf={float(c_inf(1800)):.6f} kg/kg")
    print(f"已保存：{FULL_RESULT.name}、{TABLE_RESULT.name}")


if __name__ == "__main__":
    main()
